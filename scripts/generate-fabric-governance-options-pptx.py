"""Generate a 3-slide HaadThip Fabric governance architecture deck.

Visual fingerprint:
- palette: pale mint, HaadThip green/red, deep slate, blue identity accent
- image_style: diagram-led, native editable PowerPoint shapes, no stock imagery
- silhouette: title/header + trust-zone architecture path + decision footer
- motif: 2-tone top strip, numbered control gates, green approved path, red risk flag
- forbid: dark-only canvas, generic equal cards, diagonal connectors, decorative shadows
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "presentations" / "fabric-governance-3-options.pptx"
LOGO = ROOT.parent / "presentation" / "haadthip_logo_perfect.png"

C = {
    "bg": "F4F7F5", "white": "FFFFFF", "green": "006C43", "green2": "DDF3E8",
    "red": "E60000", "red2": "FDE7E7", "slate": "0B0F17", "muted": "526173",
    "line": "C9D5D0", "blue": "2563EB", "blue2": "E8F0FF", "amber": "B76E00",
    "amber2": "FFF3D6", "cyan": "008A99", "cyan2": "DFF6F8", "grey": "EAF0ED",
}
FONT = "Aptos"
THAI = "Sarabun"


def rgb(h): return RGBColor.from_string(h)

def set_run_font(run, size, bold=False, color="slate", face=FONT):
    run.font.name = face; run.font.size = Pt(size); run.font.bold = bold
    run.font.color.rgb = rgb(C.get(color, color))
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        old = rPr.find(parse_xml(f"<{tag} {nsdecls('a')} typeface='{face}'/>" ).tag)
        if old is not None: rPr.remove(old)
        rPr.append(parse_xml(f"<{tag} {nsdecls('a')} typeface='{face}'/>"))


def textbox(slide, x, y, w, h, text, size=12, bold=False, color="slate", align=PP_ALIGN.LEFT,
            face=THAI, margin=0.02, valign=MSO_ANCHOR.MIDDLE):
    sh = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = sh.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin); tf.margin_top = tf.margin_bottom = Inches(0)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text; set_run_font(r, size, bold, color, face)
    return sh


def box(slide, x, y, w, h, fill="white", line="line", radius=True, lw=1.0):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    sh = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = rgb(C.get(fill, fill))
    sh.line.color.rgb = rgb(C.get(line, line)); sh.line.width = Pt(lw)
    if radius:
        try: sh.adjustments[0] = 0.08
        except Exception: pass
    return sh


def pill(slide, x, y, w, text, fill, color="white", size=8.5):
    sh = box(slide, x, y, w, .28, fill=fill, line=fill, radius=True, lw=.5)
    textbox(slide, x+.04, y+.01, w-.08, .24, text, size=size, bold=True, color=color,
            align=PP_ALIGN.CENTER, face=FONT)
    return sh


def arrow(slide, x1, y, x2, color="green", dashed=False):
    """Draw a short directional connector with a visible arrowhead."""
    c = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y), Inches(x2-.10), Inches(y)
    )
    c.line.color.rgb = rgb(C[color]); c.line.width = Pt(2.2)
    if dashed: c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    tip = slide.shapes.add_shape(
        MSO_SHAPE.ISOSCELES_TRIANGLE, Inches(x2-.14), Inches(y-.075), Inches(.16), Inches(.15)
    )
    tip.rotation = 90
    tip.fill.solid(); tip.fill.fore_color.rgb = rgb(C[color]); tip.line.fill.background()
    return c


def zone(slide, x, y, w, h, label, accent="green"):
    sh = box(slide, x, y, w, h, fill="white", line=accent, radius=True, lw=1.0)
    sh.fill.transparency = 6
    pill(slide, x+.15, y+.12, 1.45, label, accent, "white", 7.5)
    return sh


def icon_circle(slide, x, y, label, fill="green", text_color="white"):
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(.42), Inches(.42))
    sh.fill.solid(); sh.fill.fore_color.rgb = rgb(C[fill]); sh.line.fill.background()
    textbox(slide, x, y+.005, .42, .39, label, 10, True, text_color, PP_ALIGN.CENTER, FONT)


def node(slide, x, y, w, h, step, tag, title, detail, fill="white", line="line", badge="green"):
    box(slide, x, y, w, h, fill=fill, line=line, radius=True, lw=1.15)
    icon_circle(slide, x+.14, y+.16, str(step), badge)
    pill(slide, x+.66, y+.18, min(w-.82, 1.2), tag, badge, "white", 7.3)
    textbox(slide, x+.16, y+.72, w-.32, .47, title, 12.5, True, "slate", face=THAI,
            valign=MSO_ANCHOR.TOP)
    textbox(slide, x+.16, y+1.19, w-.32, h-1.31, detail, 8.8, False, "muted", face=THAI,
            valign=MSO_ANCHOR.TOP)


def header(slide, option, title, subtitle, badge_color):
    # background first
    bg = slide.background.fill; bg.solid(); bg.fore_color.rgb = rgb(C["bg"])
    topg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(9.333), Inches(.075))
    topg.fill.solid(); topg.fill.fore_color.rgb = rgb(C["green"]); topg.line.fill.background()
    topr = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(9.333), 0, Inches(4), Inches(.075))
    topr.fill.solid(); topr.fill.fore_color.rgb = rgb(C["red"]); topr.line.fill.background()
    if LOGO.exists(): slide.shapes.add_picture(str(LOGO), Inches(.48), Inches(.29), height=Inches(.42))
    pill(slide, 10.92, .32, 1.85, option, badge_color, "white", 8)
    textbox(slide, .48, .88, 12.2, .48, title, 23, True, "slate", face=THAI)
    textbox(slide, .5, 1.34, 12.1, .34, subtitle, 10.5, False, "muted", face=THAI)


def footer(slide, verdict, why, risk, color):
    box(slide, .48, 6.53, 12.37, .58, fill="white", line=color, radius=True, lw=1.1)
    pill(slide, .63, 6.67, 1.28, verdict, color, "white", 8)
    textbox(slide, 2.08, 6.61, 6.62, .4, why, 9.3, True, "slate", face=THAI)
    textbox(slide, 8.88, 6.61, 3.75, .4, "ข้อควรระวัง: " + risk, 8.3, False, "red", face=THAI)
    textbox(slide, .5, 7.22, 12.3, .18, "GENIE HAADTHIP CHAT  •  FABRIC DATA ACCESS GOVERNANCE  •  DIO", 6.8, True, "muted", face=FONT)


def add_slide(prs, spec):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, spec["option"], spec["title"], spec["subtitle"], spec["badge"])

    # trust zones before arrows and nodes
    zone(s, .48, 1.88, 4.05, 4.35, "GENIE / OWUI", "green")
    zone(s, 4.70, 1.88, 3.95, 4.35, spec["mid_zone"], spec["mid_color"])
    zone(s, 8.82, 1.88, 4.03, 4.35, "MICROSOFT FABRIC", "cyan")

    xs = [.72, 2.60, 4.94, 6.82, 9.06, 10.94]
    widths = [1.55, 1.55, 1.55, 1.55, 1.55, 1.55]
    y = 2.56; h = 2.75
    # connectors first, colored by the destination trust zone
    connector_colors = ["green", spec["mid_color"], spec["mid_color"], "cyan", "cyan"]
    for i in range(5):
        arrow(s, xs[i]+widths[i]+.05, 3.96, xs[i+1]-.08, color=connector_colors[i])
    for idx, n in enumerate(spec["nodes"]):
        node(s, xs[idx], y, widths[idx], h, idx+1, n[0], n[1], n[2], n[3], n[4], n[5])

    # semantic labels + one clearly bounded cross-zone control strip
    textbox(s, .77, 5.43, 3.5, .30, "① จำกัด Skill/Tool ก่อนเห็นข้อมูลจริง", 9.2, True, "green", face=THAI)
    textbox(s, 4.99, 5.43, 3.4, .30, spec["mid_note"], 9.2, True, spec["mid_color"], face=THAI)
    textbox(s, 9.11, 5.43, 3.4, .30, "③ Fabric ตัดสินสิทธิ์ข้อมูล", 9.2, True, "cyan", face=THAI)
    box(s, .72, 5.78, 11.89, .32, fill="grey", line="line", radius=True, lw=.6)
    textbox(s, .86, 5.81, 11.58, .22, spec["control_line"], 8.2, False, "muted",
            align=PP_ALIGN.CENTER, face=THAI)
    footer(s, spec["verdict"], spec["why"], spec["risk"], spec["badge"])
    return s


def build():
    prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    specs = [
        {
            "option":"OPTION A  •  TARGET", "badge":"green", "mid_zone":"DELEGATED IDENTITY", "mid_color":"blue",
            "title":"Delegated User Token / On-Behalf-Of",
            "subtitle":"สิทธิ์ของผู้ถามเดินทางถึง Fabric โดยตรง — เหมาะเป็น Production Target",
            "nodes":[
                ("USER","Internal User","Entra ID sign-in\nOID + Groups","white","green","green"),
                ("GATE","OWUI Skill Access","แสดงเฉพาะ Skill/Tool\nตาม OWUI Group","green2","green","green"),
                ("OBO","Token Broker","แลก token แบบ\nOn-Behalf-Of","blue2","blue","blue"),
                ("GUARD","Query Guard","ตรวจ Table/Column\nSELECT-only + Limit","white","blue","blue"),
                ("SQL","Fabric SQL Endpoint","Execute ด้วย\nUser Token","cyan2","cyan","cyan"),
                ("POLICY","Native Controls","SQL Role • RLS\nColumn • Masking","white","cyan","cyan"),
            ],
            "mid_note":"② Identity ผู้ใช้จริงอยู่ตลอดเส้นทาง",
            "control_line":"Audit ระบุ End User ได้  •  Catalog discovery และ Query ใช้ principal เดียวกัน  •  Deny-by-default เมื่อ token ใช้ไม่ได้",
            "verdict":"แนะนำที่สุด", "why":"Security boundary อยู่ที่ Fabric และ audit ผูกกลับถึงผู้ใช้จริง", "risk":"ต้องปรับ SSO/OBO และดูแล token lifecycle"
        },
        {
            "option":"OPTION B  •  MVP", "badge":"amber", "mid_zone":"AUTHORIZATION BROKER", "mid_color":"amber",
            "title":"Authorization Broker + Service Principal",
            "subtitle":"ทำ POC ได้เร็วบนโครงสร้างปัจจุบัน — ต้องปิด Generic SQL และใช้ Curated Views",
            "nodes":[
                ("USER","Internal User","Entra ID sign-in\nOID + Groups","white","green","green"),
                ("GATE","OWUI Skill Access","แสดงเฉพาะ Skill/Tool\nตาม OWUI Group","green2","green","green"),
                ("POLICY","Policy Broker","Group → View\nColumn → Row Scope","amber2","amber","amber"),
                ("API","Typed Data Tools","Allowlist + Inject filter\nห้าม Model ระบุ role","white","amber","amber"),
                ("SPN","Genie Service Principal","สิทธิ์เฉพาะ\nCurated Views","cyan2","cyan","cyan"),
                ("AUDIT","Fabric + App Audit","Fabric เห็น SPN\nApp log เก็บ OID","white","cyan","cyan"),
            ],
            "mid_note":"② Broker enforce ทุก Request — ไม่เชื่อ LLM",
            "control_line":"ปิด query_fabric(sql) สำหรับผู้ใช้ทั่วไป  •  Tool แบบ typed เท่านั้น  •  Service Principal อ่านเฉพาะ curated views",
            "verdict":"เหมาะกับ MVP", "why":"เปลี่ยนน้อยและส่งมอบเร็ว แต่ authorization logic อยู่ที่แอป", "risk":"Tool bug อาจใช้สิทธิ์รวมของ Service Principal"
        },
        {
            "option":"OPTION C  •  TRANSITION", "badge":"red", "mid_zone":"ROLE ROUTER", "mid_color":"red",
            "title":"Service Principal แยกตาม Role",
            "subtitle":"แยก Blast Radius ต่อ Domain แต่เพิ่มภาระ Identity และไม่เห็น End User ใน Fabric",
            "nodes":[
                ("USER","Internal User","Entra ID sign-in\nOID + Groups","white","green","green"),
                ("GATE","OWUI Skill Access","แสดงเฉพาะ Skill/Tool\nตาม OWUI Group","green2","green","green"),
                ("ROUTE","Role Router","Sales / Finance / HR\nเลือก identity","red2","red","red"),
                ("VAULT","Identity Pool","SPN ต่อ Role\nManaged Identity/Secret","white","red","red"),
                ("SQL","Fabric SQL Endpoint","Execute ด้วย\nRole SPN","cyan2","cyan","cyan"),
                ("ROLE","Role-bound Views","GRANT ต่อ Domain\nAudit เห็น Role SPN","white","cyan","cyan"),
            ],
            "mid_note":"② แยก Identity ตาม Role เพื่อลด Blast Radius",
            "control_line":"Group mapping ต้อง deterministic  •  ห้าม fallback ไป SPN สิทธิ์สูงกว่า  •  Rotation / ownership / audit mapping ต้องมีเจ้าของ",
            "verdict":"ใช้ชั่วคราว", "why":"ดีกว่า SPN เดียว แต่ role combinations และ operational overhead โตเร็ว", "risk":"Audit ไม่เห็น End User และ RLS รายบุคคลยาก"
        },
    ]
    for spec in specs: add_slide(prs, spec)
    OUT.parent.mkdir(parents=True, exist_ok=True); prs.save(OUT)
    print(OUT)

if __name__ == "__main__": build()
