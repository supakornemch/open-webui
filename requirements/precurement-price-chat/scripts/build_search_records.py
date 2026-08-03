#!/usr/bin/env python3
"""Build Azure AI Search records from the normalized Procurement template v2.

The source workbook is relational: product, variant, vendor, and vendor price
data live in separate sheets. This exporter joins those sheets itself because
Python readers do not recalculate Excel VLOOKUP formulas.

Only rows explicitly marked both ``Is Winner`` and ``Meets Spec`` are emitted.
It is intentionally an error to use a template that omits those controls: a
lowest quote is not necessarily an acceptable awarded quote.
"""

import argparse
import json
import re
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = ROOT / "scripts/Procurement_Price_Template_v2.xlsx"
DEFAULT_OUTPUT = ROOT / "procurement_prices_for_search.jsonl"
REQUIRED_SHEETS = {"Products", "Variants", "Vendors", "Vendor_Prices"}
REQUIRED_PRICE_COLUMNS = {
    "Product Code",
    "Variant Code",
    "Vendor Code",
    "Price",
    "Unit",
    "Qty",
    "Qty Range",
    "Year",
    "Is Winner",
    "Meets Spec",
    "Notes",
}


def read_rows(workbook: Any, sheet_name: str) -> list[dict[str, Any]]:
    sheet = workbook[sheet_name]
    headers = [str(cell.value or "").strip() for cell in next(sheet.iter_rows(max_row=1))]
    rows = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        if not any(value is not None and str(value).strip() for value in values):
            continue
        rows.append({header: values[index] if index < len(values) else None for index, header in enumerate(headers)})
    return rows


def as_bool(value: Any) -> bool:
    return str(value).strip().upper() in {"TRUE", "YES", "1"}


def compact_id(value: str) -> str:
    return re.sub(r"[^A-Za-z0-9_-]+", "_", value).strip("_")


def display(value: Any) -> str:
    return "" if value is None else str(value).strip()


def build_records(source: Path) -> tuple[list[dict[str, Any]], list[str]]:
    workbook = load_workbook(source, read_only=True, data_only=False)
    missing_sheets = REQUIRED_SHEETS - set(workbook.sheetnames)
    if missing_sheets:
        raise ValueError(f"Missing sheets: {', '.join(sorted(missing_sheets))}")

    products = read_rows(workbook, "Products")
    variants = read_rows(workbook, "Variants")
    vendors = read_rows(workbook, "Vendors")
    prices = read_rows(workbook, "Vendor_Prices")
    workbook.close()

    price_columns = set(prices[0]) if prices else set()
    missing_columns = REQUIRED_PRICE_COLUMNS - price_columns
    if missing_columns:
        raise ValueError(
            "Vendor_Prices must include awarded-quote controls. Missing: "
            + ", ".join(sorted(missing_columns))
        )

    product_by_code = {display(row["Product Code"]): row for row in products}
    variant_by_code = {display(row["Variant Code"]): row for row in variants}
    vendor_by_code = {display(row["Vendor Code"]): row for row in vendors}
    records = []
    missing_references = []
    ignored_not_awarded = 0

    for row in prices:
        if not as_bool(row["Is Winner"]) or not as_bool(row["Meets Spec"]):
            ignored_not_awarded += 1
            continue

        product_code = display(row["Product Code"])
        variant_code = display(row["Variant Code"])
        vendor_code = display(row["Vendor Code"])
        product = product_by_code.get(product_code)
        variant = variant_by_code.get(variant_code) if variant_code else None
        vendor = vendor_by_code.get(vendor_code)
        if not product or not vendor or (variant_code and not variant):
            missing_references.append(
                f"product={product_code!r}, variant={variant_code!r}, vendor={vendor_code!r}"
            )
            continue

        price = row["Price"]
        if not isinstance(price, (int, float)) or price < 0:
            raise ValueError(f"Invalid Price for product {product_code}: {price!r}")

        item_name = display(product.get("Product Name"))
        variant_description = display(variant.get("Description")) if variant else ""
        vendor_name = display(vendor.get("Vendor Name"))
        year = int(row["Year"])
        quantity_range = display(row["Qty Range"])
        source_key = "|".join((product_code, variant_code or "base", vendor_code, str(year), quantity_range))
        record = {
            "id": compact_id(source_key),
            "product_code": product_code,
            "product_name": item_name,
            "category": display(product.get("Category")),
            "pricing_model": display(product.get("Pricing Model")),
            "variant_code": variant_code,
            "variant_description": variant_description,
            "vendor_code": vendor_code,
            "vendor_name": vendor_name,
            "year": year,
            "price": float(price),
            "unit": display(row["Unit"]),
            "qty": row["Qty"],
            "qty_range": quantity_range,
            "is_winner": True,
            "meets_spec": True,
            "notes": display(row["Notes"]),
            "corpus": "procurement-prices",
            "source": source.name,
        }
        product_text = " ".join(part for part in (item_name, variant_description) if part)
        record["content"] = (
            f"{product_text} | หมวด {record['category']} | ราคา {record['price']:g} บาท/{record['unit']} "
            f"| ช่วงจำนวน {quantity_range or 'ไม่ระบุ'} | ผู้ชนะ {vendor_name} | ปี {year}"
        )
        records.append(record)

    if missing_references:
        raise ValueError("Unresolved template references: " + "; ".join(missing_references[:5]))
    if not records:
        raise ValueError("No awarded rows found: mark Is Winner and Meets Spec as TRUE")

    warnings = [f"Ignored {ignored_not_awarded} non-awarded or non-compliant vendor quote(s)"]
    return records, warnings


def main() -> None:
    parser = argparse.ArgumentParser(description="Build Procurement v2 Azure AI Search records")
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    records, warnings = build_records(args.source)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("w", encoding="utf-8") as output:
        for record in records:
            output.write(json.dumps(record, ensure_ascii=False) + "\n")

    categories = sorted({record["category"] for record in records})
    print(f"Wrote {len(records)} awarded price record(s) to {args.output}")
    print(f"Categories: {', '.join(categories)}")
    for warning in warnings:
        print(f"Warning: {warning}")


if __name__ == "__main__":
    main()