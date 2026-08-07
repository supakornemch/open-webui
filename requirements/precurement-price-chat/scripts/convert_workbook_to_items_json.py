#!/usr/bin/env python3
"""Convert the legacy workbook into one JSON object per source price row.

Merged item and description cells are forward-filled. Rows that represent
quantity tiers become independent objects with the parent item values copied
into each row, so merged cells cannot shift values into the wrong column.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = ROOT / "data/source/Trade Marketing Materials price for Y2026.final.xlsx"
DEFAULT_OUTPUT = ROOT / "data/exports/Trade Marketing Materials price for Y2026.items.json"

EXCEL_ERROR_VALUES = {
    "#REF!",
    "#N/A",
    "#VALUE!",
    "#DIV/0!",
    "#NAME?",
    "#NUM!",
    "#NULL!",
}

SHEET_CONFIG = {
    "1.POSM(MKT)": {
        "category": "POSM",
        "header_row": 2,
        "item_header": ("Item",),
        "description_header": ("รายการ",),
        "quantity_header": ("QTY",),
    },
    "2.Printing(MKT)": {
        "category": "Printing-MKT",
        "header_row": 2,
        "item_header": ("Item",),
        "description_header": ("รายการ", "รายากร"),
        "quantity_header": ("QTY",),
    },
    "3.Garment": {
        "category": "Garment",
        "header_row": 2,
        "item_header": ("Item",),
        "description_header": ("รายการ",),
        "quantity_header": ("QTY Y2026",),
    },
    "4.สรุปPremium": {
        "category": "Premium",
        "header_row": 2,
        "item_header": ("Item",),
        "description_header": ("รายการ",),
        "quantity_header": ("QTY Y2026",),
    },
    "5.Printing-Rate1-43": {
        "category": "Printing-Rate",
        "header_row": 2,
        "item_header": ("ลำดับ",),
        "description_header": ("รายการ",),
        "quantity_header": ("QTY Y2026",),
    },
    "5.Printing-Rate44-109": {
        "category": "Printing-Rate",
        "header_row": 5,
        "item_header": ("ลำดับ",),
        "description_header": ("รายการ",),
        "quantity_header": ("QTY Y2026",),
    },
}


def clean_text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, float) and math.isnan(value):
        return None
    text = str(value).replace("\xa0", " ").replace("\r\n", "\n").replace("\r", "\n")
    parts = [re.sub(r"[ \t]+", " ", part).strip() for part in text.split("\n")]
    text = " ".join(part for part in parts if part)
    return text or None


def scalar(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, float):
        if math.isnan(value):
            return None
        if value.is_integer():
            return int(value)
        return value
    if isinstance(value, (int, bool)):
        return value
    return clean_text(value)


def present(value: Any) -> bool:
    normalized = scalar(value)
    return normalized is not None and not (
        isinstance(normalized, str) and normalized.strip().upper() in EXCEL_ERROR_VALUES
    )


def normalized_header(value: Any) -> str:
    return clean_text(value) or ""


def find_header_column(headers: dict[int, str], candidates: tuple[str, ...]) -> int:
    wanted = {candidate.casefold() for candidate in candidates}
    for column, header in headers.items():
        if header.casefold() in wanted:
            return column
    raise ValueError(f"Missing header {candidates!r}; available={list(headers.values())!r}")


def quantity_details(value: Any) -> dict[str, Any]:
    original = scalar(value)
    if not present(value):
        return {"value": None, "min": None, "max": None, "is_exact": False}
    text = clean_text(value)
    if text is None:
        return {"value": None, "min": None, "max": None, "is_exact": False}

    number_texts = re.findall(r"\d[\d,]*", text.replace(" ", ""))
    numbers = [int(number_text.replace(",", "")) for number_text in number_texts]
    if not numbers:
        return {"value": original, "min": None, "max": None, "is_exact": False}

    is_range = len(numbers) >= 2 and re.search(r"[-–—ถึง]|\bto\b", text, re.IGNORECASE)
    is_open = bool(re.search(r"\+|ขึ้นไป|มากกว่า|above|over", text, re.IGNORECASE))
    if is_range:
        lower, upper = numbers[0], numbers[1]
        return {
            "value": original,
            "min": lower,
            "max": upper,
            "is_exact": False,
        }
    if is_open:
        return {"value": original, "min": numbers[0], "max": None, "is_exact": False}
    if len(numbers) == 1:
        return {"value": original, "min": numbers[0], "max": numbers[0], "is_exact": True}
    return {"value": original, "min": numbers[0], "max": None, "is_exact": False}


def item_number(value: Any) -> str | None:
    normalized = scalar(value)
    if normalized is None:
        return None
    if isinstance(normalized, int):
        return str(normalized)
    if isinstance(normalized, float) and normalized.is_integer():
        return str(int(normalized))
    text = str(normalized).strip()
    if re.fullmatch(r"\d+\.0", text):
        return text[:-2]
    return text


def raw_cells(ws: Any, row_number: int) -> dict[str, Any]:
    result = {}
    for column in range(1, ws.max_column + 1):
        value = scalar(ws.cell(row_number, column).value)
        if present(value):
            result[get_column_letter(column)] = value
    return result


def row_has_data(ws: Any, row_number: int, columns: list[int]) -> bool:
    return any(present(ws.cell(row_number, column).value) for column in columns)


def vendor_columns(headers: dict[int, str]) -> list[tuple[int, str]]:
    columns = []
    for column, header in headers.items():
        match = re.fullmatch(r"vendor\s*(\d+)", header, re.IGNORECASE)
        if match:
            columns.append((column, f"Vendor {int(match.group(1))}"))
    return sorted(columns, key=lambda pair: int(pair[1].split()[-1]))


def award_columns(headers: dict[int, str], phrase: str) -> list[int]:
    return [column for column, header in headers.items() if phrase in header]


def history_values(ws: Any, headers: dict[int, str], row_number: int) -> dict[str, dict[str, Any]]:
    result = {}
    for year in (2023, 2024, 2025):
        history_header = f"History {year}".casefold()
        history_column = next(
            (column for column, header in headers.items() if header.casefold() == history_header),
            None,
        )
        if history_column is None:
            continue
        price_column = history_column + 1
        result[str(year)] = {
            "qty": scalar(ws.cell(row_number, history_column).value),
            "price": scalar(ws.cell(row_number, price_column).value),
        }
    return result


def build_tier(
    ws: Any,
    headers: dict[int, str],
    row_number: int,
    quantity_column: int,
    vendor_price_columns: list[tuple[int, str]],
    awarded_price_columns: list[int],
    awarded_vendor_columns: list[int],
) -> dict[str, Any]:
    quantity = quantity_details(ws.cell(row_number, quantity_column).value)
    prices = []
    for column, vendor in vendor_price_columns:
        price = scalar(ws.cell(row_number, column).value)
        if present(price):
            prices.append({"vendor": vendor, "price": price})

    awarded_price = None
    for column in awarded_price_columns:
        value = scalar(ws.cell(row_number, column).value)
        if present(value):
            awarded_price = value
            break

    awarded_vendor = None
    award_notes = []
    for index, column in enumerate(awarded_vendor_columns):
        value = scalar(ws.cell(row_number, column).value)
        if not present(value):
            continue
        if awarded_vendor is None:
            awarded_vendor = value
        else:
            award_notes.append(value)

    return {
        "source_row": row_number,
        "quantity": quantity,
        "vendor_prices": prices,
        "awarded_price": awarded_price,
        "awarded_vendor": awarded_vendor,
        "award_notes": award_notes,
        "raw": raw_cells(ws, row_number),
    }


def build_sheet_items(ws: Any, config: dict[str, Any]) -> list[dict[str, Any]]:
    header_row = config["header_row"]
    headers = {
        column: normalized_header(ws.cell(header_row, column).value)
        for column in range(1, ws.max_column + 1)
        if normalized_header(ws.cell(header_row, column).value)
    }
    item_column = find_header_column(headers, config["item_header"])
    description_column = find_header_column(headers, config["description_header"])
    quantity_column = find_header_column(headers, config["quantity_header"])
    vendors = vendor_columns(headers)
    awarded_price_columns = award_columns(headers, "ราคาที่ผ่านการประมูล")
    awarded_vendor_columns = award_columns(headers, "รายชื่อผู้ผ่านการประมูล")
    row_data_columns = [quantity_column]
    row_data_columns.extend(column for column, _ in vendors)
    row_data_columns.extend(awarded_price_columns)
    row_data_columns.extend(awarded_vendor_columns)

    columns = [
        {"column": get_column_letter(column), "header": header or None}
        for column, header in (
            (column, normalized_header(ws.cell(header_row, column).value))
            for column in range(1, ws.max_column + 1)
        )
    ]
    items = []
    current = None

    def finish_current() -> None:
        if current is None:
            return
        if not current["quantity_tiers"]:
            return
        current["name"] = " ".join(current.pop("_description_parts")) or None
        current["description_lines"] = current.pop("_description_lines")
        current["tier_count"] = len(current["quantity_tiers"])
        current["source"]["row_end"] = current["quantity_tiers"][-1]["source_row"] if current["quantity_tiers"] else current["source"]["row_start"]
        items.append(current)

    for row_number in range(header_row + 1, ws.max_row + 1):
        raw_item = item_number(ws.cell(row_number, item_column).value)
        raw_description = clean_text(ws.cell(row_number, description_column).value)
        has_row_data = row_has_data(ws, row_number, row_data_columns)
        if raw_item is not None:
            finish_current()
            current = {
                "item_number": raw_item,
                "name": None,
                "description_lines": [],
                "category": config["category"],
                "quantity_tiers": [],
                "tier_count": 0,
                "history": history_values(ws, headers, row_number),
                "source": {
                    "sheet": ws.title,
                    "header_row": header_row,
                    "row_start": row_number,
                    "item_cell": f"{get_column_letter(item_column)}{row_number}",
                    "description_cell": f"{get_column_letter(description_column)}{row_number}",
                    "quantity_field": headers[quantity_column],
                    "item_column": get_column_letter(item_column),
                    "description_column": get_column_letter(description_column),
                    "quantity_column": get_column_letter(quantity_column),
                    "columns": columns,
                },
                "_description_parts": [],
                "_description_lines": [],
            }

        if current is None:
            continue
        if raw_description and (raw_item is not None or has_row_data) and raw_description not in current["_description_lines"]:
            current["_description_lines"].append(raw_description)
            current["_description_parts"].append(raw_description)

        if has_row_data:
            current["quantity_tiers"].append(
                build_tier(
                    ws,
                    headers,
                    row_number,
                    quantity_column,
                    vendors,
                    awarded_price_columns,
                    awarded_vendor_columns,
                )
            )

    finish_current()
    return items


def flatten_items(logical_items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows = []
    for logical_item in logical_items:
        source = logical_item["source"]
        for tier in logical_item["quantity_tiers"]:
            row_number = tier["source_row"]
            raw_description = tier["raw"].get(source["description_column"])
            rows.append(
                {
                    "id": f"{source['sheet']}:{logical_item['item_number']}:{row_number}",
                    "item_number": logical_item["item_number"],
                    "name": logical_item["name"],
                    "description": " ".join(logical_item["description_lines"]) or None,
                    "row_description": raw_description,
                    "description_lines": logical_item["description_lines"],
                    "category": logical_item["category"],
                    "history": logical_item["history"],
                    "quantity": tier["quantity"],
                    "vendor_prices": tier["vendor_prices"],
                    "awarded_price": tier["awarded_price"],
                    "awarded_vendor": tier["awarded_vendor"],
                    "award_notes": tier["award_notes"],
                    "source": {
                        "sheet": source["sheet"],
                        "header_row": source["header_row"],
                        "row": row_number,
                        "item_row": source["row_start"],
                        "item_cell": f"{source['item_column']}{source['row_start']}",
                        "description_cell": f"{source['description_column']}{row_number}",
                        "quantity_cell": f"{source['quantity_column']}{row_number}",
                        "quantity_field": source["quantity_field"],
                    },
                    "raw": tier["raw"],
                }
            )
    return rows


def sheet_summary(ws: Any) -> dict[str, Any]:
    non_empty = sum(
        1
        for row in ws.iter_rows()
        for cell in row
        if scalar(cell.value) is not None
    )
    return {
        "name": ws.title,
        "max_row": ws.max_row,
        "max_column": ws.max_column,
        "non_empty_cell_count": non_empty,
        "merged_ranges": [str(value) for value in ws.merged_cells.ranges],
    }


def metadata_sheet(ws: Any) -> dict[str, Any]:
    rows = []
    for row_number in range(1, ws.max_row + 1):
        values = raw_cells(ws, row_number)
        if values:
            rows.append({"source_row": row_number, "cells": values})
    return {"name": ws.title, "reason": "metadata-only sheet", "rows": rows}


def convert(source: Path) -> dict[str, Any]:
    workbook = openpyxl.load_workbook(source, data_only=True, read_only=False)
    logical_items = []
    summaries = []
    non_item_sheets = []
    for ws in workbook.worksheets:
        summaries.append(sheet_summary(ws))
        config = SHEET_CONFIG.get(ws.title)
        if config is None:
            non_item_sheets.append(metadata_sheet(ws))
            continue
        logical_items.extend(build_sheet_items(ws, config))

    items = flatten_items(logical_items)

    return {
        "source_file": source.name,
        "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        "format": "one source price row per object; merged item values forward-filled",
        "sheet_count": len(workbook.sheetnames),
        "item_count": len(items),
        "logical_item_count": len(logical_items),
        "quantity_tier_count": len(items),
        "sheet_summaries": summaries,
        "non_item_sheets": non_item_sheets,
        "items": items,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    if not args.source.exists():
        raise FileNotFoundError(args.source)

    payload = convert(args.source)
    args.output.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2, allow_nan=False) + "\n",
        encoding="utf-8",
    )
    print(f"Wrote {args.output}: {args.output.stat().st_size} bytes")
    print(f"Rows/items: {payload['item_count']}")
    print(f"Logical items: {payload['logical_item_count']}")
    for summary in payload["sheet_summaries"]:
        count = sum(1 for item in payload["items"] if item["source"]["sheet"] == summary["name"])
        print(f"- {summary['name']}: {count} item(s)")


if __name__ == "__main__":
    main()