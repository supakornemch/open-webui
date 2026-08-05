#!/usr/bin/env python3
"""
Extract product catalog from Excel to JSON format for LLM chatbot.

Output structure:
{
  "items": [
    {
      "item_id": "POSM-5",
      "category": "POSM",
      "item_number": "5",
      "base_name": "ร่มโค้ก 36 นิ้ว",
      "full_name": "ร่มโค้ก 36 นิ้ว - โครงเหล็ก สีตาย (ไม่เกิน 4 สี)",
      "specs": {
        "ขนาด": "36 นิ้ว",
        "โครง": "เหล็ก",
        "สี": "สีตาย"
      },
      "vendor_prices": [...],
      "winner_price": 523,
      "winner_vendor": "Vendor 9",
      "price_range": {"min": 523, "max": 1950},
      "vendor_count": 4
    }
  ],
  "metadata": {
    "total_items": 201,
    "categories": ["POSM", "Printing", "Garment", "Premium", "PrintRate"],
    "extracted_at": "2026-07-31T..."
  }
}
"""

import openpyxl
import json
from datetime import datetime
from pathlib import Path
import re

EXCEL_FILE = "Trade Marketing Materials price for Y2026.final.xlsx"
OUTPUT_FILE = "data/catalog.json"

CATEGORY_MAP = {
    '1.POSM(MKT)': 'POSM',
    '2.Printing(MKT)': 'Printing',
    '3.Garment': 'Garment',
    '4.สรุปPremium': 'Premium',
    '5.Printing-Rate1-43': 'PrintRate',
    '5.Printing-Rate44-109': 'PrintRate'
}

def extract_base_name(full_name):
    """Extract base product name before specs/modifiers."""
    # Remove everything after: -,  (,  \n
    patterns = [
        r'^([^-\n(]+)',  # Before -, \n, or (
    ]
    
    for pattern in patterns:
        match = re.match(pattern, full_name)
        if match:
            return match.group(1).strip()
    
    return full_name.split()[0] if full_name else full_name

def extract_specs(full_name):
    """Extract specifications from full name."""
    specs = {}
    
    # ขนาด
    size_match = re.search(r'(\d+)\s*นิ้ว', full_name)
    if size_match:
        specs['ขนาด'] = f"{size_match.group(1)} นิ้ว"
    
    size_match = re.search(r'(\d+x\d+(?:x\d+)?)\s*(?:cm|ซม)', full_name, re.IGNORECASE)
    if size_match:
        specs['ขนาด'] = size_match.group(1)
    
    # โครง
    if 'โครงเหล็ก' in full_name:
        specs['โครง'] = 'เหล็ก'
    elif 'โครงไฟเบอร์' in full_name:
        specs['โครง'] = 'ไฟเบอร์'
    
    # สี
    if 'สีตาย' in full_name:
        specs['สี'] = 'สีตาย'
    elif 'ไล่ระดับสี' in full_name:
        specs['สี'] = 'ไล่ระดับสี'
    
    # ผ้า (for Garment)
    fabric_match = re.search(r'ผ้า\s*([^\n]+?)(?=\n|$)', full_name)
    if fabric_match:
        specs['ผ้า'] = fabric_match.group(1).strip()
    
    return specs

def extract_vendor_prices(row, start_col=10, max_vendors=14):
    """Extract vendor prices from row (columns K-X, indices 10-23)."""
    prices = []
    
    for i in range(start_col, start_col + max_vendors):
        if i >= len(row):
            break
        
        value = row[i]
        if value and isinstance(value, (int, float)) and value > 0:
            vendor_num = i - start_col + 1
            prices.append({
                'vendor': f"Vendor {vendor_num}",
                'price': float(value)
            })
    
    return prices

def extract_catalog():
    """Extract catalog from all Excel sheets."""
    wb = openpyxl.load_workbook(EXCEL_FILE)
    
    catalog = []
    
    for sheet_name in wb.sheetnames:
        if sheet_name not in CATEGORY_MAP:
            continue
        
        category = CATEGORY_MAP[sheet_name]
        ws = wb[sheet_name]
        
        print(f"Processing {sheet_name} ({category})...")
        
        for row in ws.iter_rows(min_row=3, values_only=True):
            # Skip if no item number or name
            if not row[0] or not row[2]:
                continue
            
            item_num = str(row[0])
            full_name = str(row[2])
            
            # Skip if name is too short (likely header or junk)
            if len(full_name) < 5:
                continue
            
            # Extract prices
            vendor_prices = extract_vendor_prices(row)
            
            if not vendor_prices:
                print(f"  ⚠️  Skipping {item_num}: No vendor prices")
                continue
            
            # Calculate winner
            prices_only = [p['price'] for p in vendor_prices]
            winner_price = min(prices_only)
            winner_vendor = next(p['vendor'] for p in vendor_prices if p['price'] == winner_price)
            
            # Generate item_id
            item_id = f"{category}-{item_num}"
            
            # Extract base name and specs
            base_name = extract_base_name(full_name)
            specs = extract_specs(full_name)
            
            item = {
                'item_id': item_id,
                'category': category,
                'item_number': item_num,
                'base_name': base_name,
                'full_name': full_name,
                'specs': specs,
                'vendor_prices': vendor_prices,
                'winner_price': winner_price,
                'winner_vendor': winner_vendor,
                'price_range': {
                    'min': min(prices_only),
                    'max': max(prices_only)
                },
                'vendor_count': len(vendor_prices)
            }
            
            catalog.append(item)
    
    wb.close()
    
    # Create output
    output = {
        'items': catalog,
        'metadata': {
            'total_items': len(catalog),
            'categories': list(set(item['category'] for item in catalog)),
            'extracted_at': datetime.now().isoformat(),
            'source_file': EXCEL_FILE
        }
    }
    
    return output

def main():
    print("=== Extracting Product Catalog ===\n")
    
    # Create output directory
    Path("data").mkdir(exist_ok=True)
    
    # Extract
    catalog = extract_catalog()
    
    # Save
    with open(OUTPUT_FILE, 'w', encoding='utf-8') as f:
        json.dump(catalog, f, ensure_ascii=False, indent=2)
    
    print(f"\n✅ Extracted {catalog['metadata']['total_items']} items")
    print(f"📁 Saved to: {OUTPUT_FILE}")
    
    # Summary by category
    print("\n=== Summary by Category ===")
    category_counts = {}
    for item in catalog['items']:
        cat = item['category']
        category_counts[cat] = category_counts.get(cat, 0) + 1
    
    for cat, count in sorted(category_counts.items()):
        print(f"  {cat}: {count} items")

if __name__ == '__main__':
    main()
