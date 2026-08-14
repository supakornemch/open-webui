#!/usr/bin/env python3
"""Export master data from the legacy workbook into the NEW minimal catalog
master format (no vendor columns).

Output columns (exactly):
   ปี | ลำดับ | รายการ | จำนวนขั้นต่ำ | จำนวนสูงสุด | การคิดราคา |
   ราคาที่ผ่านการประมูล | รายชื่อผู้ผ่านการประมูล | หมายเหตุ

Only Printing-MKT + Printing-Rate (per the current product decision).
One sheet per category. Rows that belong to the same product leave
ลำดับ/รายการ/การคิดราคา/รายชื่อผู้ผ่านการประมูล/หมายเหตุ blank on
continuation tiers (forward-filled by the ingest pipeline).

Run: python3 scripts/export_catalog_master_minimal.py
"""

from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from convert_workbook_to_items_json import SHEET_CONFIG, build_sheet_items, present
from convert_workbook_to_catalog_master import (
    SHEETS,
    VENDOR_NAME_MAP,
    normalize_award_vendor,
    parse_vendor_cell,
    split_quantity,
    vendor_column_name,
    vendor_number,
)

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "data/source/Trade Marketing Materials price for Y2026.final.xlsx"
OUTPUT = ROOT / "data/masters/Procurement-Pricing-Assistant-Master-2026_MINIMAL.xlsx"

YEAR = 2026
CATEGORIES = ["POSM", "Printing-MKT", "Printing-Rate", "Garment", "Premium-EA&HRC"]

# exact column set (no vendor columns)
HEADERS = ["ปี", "ลำดับ", "รายการ", "จำนวนขั้นต่ำ", "จำนวนสูงสุด", "การคิดราคา",
           "ราคาที่ผ่านการประมูล", "รายชื่อผู้ผ่านการประมูล", "หมายเหตุ"]

HEADER_FILL = PatternFill("solid", start_color="1F4E79", end_color="1F4E79")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=11)
BODY_FONT = Font(name="Arial", size=10)
NOTE_ALIGN = Alignment(vertical="top", wrap_text=True)
HEADER_ALIGN = Alignment(vertical="center", horizontal="center", wrap_text=True)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
RED_FILL = PatternFill("solid", start_color="FFC7CE", end_color="FFC7CE")
RED_FONT = Font(name="Arial", size=10, color="9C0006")


def export_rows() -> dict[str, list[list[object]]]:
    """Return {category: [row, ...]} where each row is the 9-column export."""
    workbook = openpyxl.load_workbook(SOURCE, data_only=True, read_only=False)
    by_cat: dict[str, list[list[object]]] = {}
    for sheet_name in SHEETS:
        config = SHEET_CONFIG[sheet_name]
        ws = workbook[sheet_name]
        category = config["category"]
        cat_rows: list[list[object]] = []
        for item in build_sheet_items(ws, config):
            first_tier = True
            for tier in item["quantity_tiers"]:
                qty_min, qty_max = split_quantity(tier["quantity"])

                # collect vendor spec/flag notes only on the first tier row
                # (product-level หมายเหตุ), from ALL tiers of the product
                if first_tier:
                    vendor_notes = []
                    for t in item["quantity_tiers"]:
                        for quote in t["vendor_prices"]:
                            number = vendor_number(quote["vendor"])
                            if not number:
                                continue
                            col_name = vendor_column_name(category, number)
                            _price, spec, flag = parse_vendor_cell(quote["price"])
                            if spec:
                                vendor_notes.append(f"[{col_name}: สเปค {spec}]")
                            if flag:
                                vendor_notes.append(f"[{col_name}: ราคา/จำนวน {flag} รอยืนยัน]")
                    notes = [str(n) for n in tier["award_notes"] if present(n)]
                    notes.extend(list(dict.fromkeys(vendor_notes)))
                    award_name = normalize_award_vendor(tier["awarded_vendor"], category)
                    note_text = "; ".join(notes) if notes else None
                else:
                    award_name = None
                    note_text = None

                cat_rows.append([
                    YEAR,
                    item["item_number"] if first_tier else None,
                    item["name"] if first_tier else None,
                    qty_min,
                    qty_max,
                    None,  # การคิดราคา (Excel Master fills)
                    tier["awarded_price"],
                    award_name or None,  # รายชื่อผู้ผ่านการประมูล (product-level)
                    note_text,
                ])
                first_tier = False
        by_cat.setdefault(category, []).extend(cat_rows)
    return by_cat


def build_workbook(rows_by_cat: dict[str, list[list[object]]]) -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    col_of = {h: i + 1 for i, h in enumerate(HEADERS)}
    last_col = len(HEADERS)

    for category in CATEGORIES:
        ws = wb.create_sheet(category)
        for col, header in enumerate(HEADERS, start=1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = HEADER_ALIGN
            cell.border = BORDER

        cat_rows = rows_by_cat.get(category, [])
        for r, row in enumerate(cat_rows, start=2):
            for c, value in enumerate(row, start=1):
                if value is None:
                    continue
                cell = ws.cell(row=r, column=c, value=value)
                cell.font = BODY_FONT
                cell.border = BORDER
                cell.alignment = (
                    NOTE_ALIGN
                    if HEADERS[c - 1] == "หมายเหตุ"
                    else Alignment(vertical="top", wrap_text=True)
                )
                if HEADERS[c - 1] in ("จำนวนขั้นต่ำ", "จำนวนสูงสุด", "ราคาที่ผ่านการประมูล"):
                    cell.number_format = "#,##0"
            for c in range(1, last_col + 1):
                ws.cell(row=r, column=c).border = BORDER

        last_row = len(cat_rows) + 1
        width_map = {
            col_of["ปี"]: 8, col_of["ลำดับ"]: 8, col_of["รายการ"]: 52,
            col_of["จำนวนขั้นต่ำ"]: 12, col_of["จำนวนสูงสุด"]: 12,
            col_of["การคิดราคา"]: 12, col_of["ราคาที่ผ่านการประมูล"]: 16,
            col_of["รายชื่อผู้ผ่านการประมูล"]: 30, col_of["หมายเหตุ"]: 42,
        }
        for col, w in width_map.items():
            ws.column_dimensions[get_column_letter(col)].width = w

        ws.freeze_panes = "G2"
        ws.auto_filter.ref = f"A1:{get_column_letter(last_col)}{last_row}"

        # conditional formatting only when there are data rows
        if last_row > 1:
            qty_min_letter = get_column_letter(col_of["จำนวนขั้นต่ำ"])
            award_letter = get_column_letter(col_of["ราคาที่ผ่านการประมูล"])
            ws.conditional_formatting.add(
                f"{award_letter}2:{award_letter}{last_row}",
                FormulaRule(formula=[f'AND({qty_min_letter}2<>"",{award_letter}2="")'],
                            fill=RED_FILL, font=RED_FONT),
            )
            year_letter = get_column_letter(col_of["ปี"])
            item_letter = get_column_letter(col_of["ลำดับ"])
            ws.conditional_formatting.add(
                f"{item_letter}2:{item_letter}{last_row}",
                FormulaRule(formula=[f'AND(${item_letter}2<>"",COUNTIFS(${year_letter}$2:${year_letter}${last_row},${year_letter}2,${item_letter}$2:${item_letter}${last_row},${item_letter}2)>1)'],
                            fill=RED_FILL, font=RED_FONT),
            )

    # Legend
    lg = wb.create_sheet("Legend")
    lg.column_dimensions["A"].width = 26
    lg.column_dimensions["D"].width = 78
    lg["A1"] = "Catalog Master MINIMAL — ไม่มีคอลัมน์ vendor (9 คอลัมน์)"
    lg["A1"].font = Font(name="Arial", bold=True, size=13, color="1F4E79")
    for col, text in enumerate(["คอลัมน์", "ชนิด", "บังคับ", "ความหมาย"], start=1):
        cell = lg.cell(row=3, column=col, value=text)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BORDER
    legend = [
        ("ปี", "number", "✔", "ปีของราคา (2025/2026)"),
        ("ลำดับ", "text", "✔", "เลข item เริ่ม 1 ใหม่ทุกหมวด — แถวแรกของ item, เรทต่อเนื่องเว้นว่าง"),
        ("รายการ", "text", "✔", "ชื่อ/สเปคสินค้า — แถวแรกของ item, เรทต่อเนื่องเว้นว่าง"),
        ("จำนวนขั้นต่ำ", "number", "✔", "ขั้นต่ำของเรท — 500 = 500+, range 1, exact 5000"),
        ("จำนวนสูงสุด", "number", "✘", "สูงสุดของเรท — ว่าง = ขึ้นไป, exact = เท่า min"),
        ("การคิดราคา", "text", "✘", "หน่วยคิดราคา เช่น ต่อชิ้น / ต่อ ตรม. (กรอกเอง)"),
        ("ราคาที่ผ่านการประมูล", "number", "✔", "ราคาต่อหน่วยที่ชนะ จาก workbook"),
        ("รายชื่อผู้ผ่านการประมูล", "text", "✘", "Vendor ที่ผ่านการประมูล (V5 / ชื่อ) — ระดับสินค้า"),
        ("หมายเหตุ", "text", "✘", "หมายเหตุ award — ระดับสินค้า"),
    ]
    for r, row in enumerate(legend, start=4):
        for c, value in enumerate(row, start=1):
            cell = lg.cell(row=r, column=c, value=value)
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = NOTE_ALIGN
    return wb


def main() -> None:
    rows_by_cat = export_rows()
    wb = build_workbook(rows_by_cat)
    wb.save(OUTPUT)
    print(f"Wrote {OUTPUT} ({OUTPUT.stat().st_size} bytes)")
    for cat in CATEGORIES:
        n = len(rows_by_cat.get(cat, []))
        print(f"  {cat}: {n} row(s)")


if __name__ == "__main__":
    main()
